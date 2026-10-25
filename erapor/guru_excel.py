import io
import re

import xlrd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .db import execute, query, transaction


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).strip()
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def _jam(value):
    if value is None or value == "":
        return 0
    if isinstance(value, (int, float)):
        return int(value)
    cleaned = _text(value).replace(",", ".")
    try:
        return int(float(cleaned))
    except ValueError:
        return 0


def _rows_from_xls(data):
    book = xlrd.open_workbook(file_contents=data)
    sheets = []
    for sheet in book.sheets():
        rows = []
        for r in range(sheet.nrows):
            rows.append([sheet.cell_value(r, c) for c in range(sheet.ncols)])
        sheets.append(rows)
    return sheets


def _rows_from_xlsx(data):
    book = load_workbook(io.BytesIO(data), data_only=True)
    sheets = []
    for sheet in book.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(list(row))
        sheets.append(rows)
    return sheets


def _split_nama(raw):
    text = _text(raw).strip(" ,")
    if "," in text:
        nama, gelar = text.split(",", 1)
        gelar = gelar.strip(" ,")
        if gelar:
            return nama.strip(), gelar
    return text, ""


def _header_col(row, label):
    label = label.upper()
    for index, value in enumerate(row):
        if label in _text(value).upper():
            return index
    return None


def parse_guru_rows(sheets):
    header_at = None
    for rows in sheets:
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                if _text(value).upper() == "NAMA GURU":
                    header_at = (rows, r, c)
                    break
            if header_at:
                break
        if header_at:
            break
    if not header_at:
        raise ValueError("Kolom NAMA GURU tidak ditemukan. Gunakan file jadwal yang sama.")

    rows, header_row, nama_col = header_at
    header = rows[header_row]
    mapel_col = _header_col(header, "MATA PELAJARAN")
    jam_col = _header_col(header, "JUMLAH JAM")
    kelas_col = _header_col(header, "MENGAJAR DI KELAS")
    if mapel_col is None:
        mapel_col = nama_col + 2
    if jam_col is None:
        jam_col = nama_col + 4
    if kelas_col is None:
        kelas_col = nama_col + 5

    teachers = []
    current = None
    for row in rows[header_row + 1 :]:
        def pick(index):
            if index is None or index >= len(row):
                return ""
            return row[index]

        if any(_text(cell).upper() == "CONTOH" for cell in row):
            continue
        nama = _text(pick(nama_col))
        mapel = _text(pick(mapel_col))
        kelas = _text(pick(kelas_col))
        jam = _jam(pick(jam_col))
        if nama.upper() in {"NAMA KELAS", "NAMA GURU"}:
            break
        if not nama and not mapel:
            continue
        item = {"mapel": mapel, "jam": jam, "kelas": kelas}
        if not nama:
            if current is not None and (mapel or kelas or jam):
                current["items"].append(item)
            continue
        nama_bersih, gelar = _split_nama(nama)
        current = {"nama": nama_bersih, "gelar_belakang": gelar, "items": []}
        if mapel or kelas or jam:
            current["items"].append(item)
        teachers.append(current)

    hasil = []
    for teacher in teachers:
        mapel_parts = []
        kelas_parts = []
        total = 0
        for item in teacher["items"]:
            total += item["jam"]
            if item["mapel"]:
                mapel_parts.append(f"{item['mapel']} ({item['jam']})" if item["jam"] else item["mapel"])
            if item["kelas"] and item["kelas"] not in kelas_parts:
                kelas_parts.append(item["kelas"])
        hasil.append(
            {
                "nama": teacher["nama"],
                "gelar_belakang": teacher["gelar_belakang"],
                "mapel_ajar": ", ".join(mapel_parts),
                "jumlah_jam": total,
                "mengajar_di": "; ".join(kelas_parts),
            }
        )
    if not hasil:
        raise ValueError("Tidak ada nama guru yang terbaca dari file.")
    return hasil


def formulir_guru():
    """Form kosong yang bisa diunggah lewat tombol Upload data guru."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Form Guru"
    judul = [
        "No",
        "NAMA GURU",
        "MATA PELAJARAN",
        "JUMLAH JAM",
        "MENGAJAR DI KELAS",
    ]
    contoh = [
        ["contoh", "Nama Lengkap, S.Pd.", "Bahasa Indonesia", 24, "7.1 - 7.3"],
        ["contoh", None, "Matematika", 12, "8.1"],
    ]
    ws.append(judul)
    for baris in contoh:
        ws.append(baris)
    isi = PatternFill("solid", fgColor="0F766E")
    kuning = PatternFill("solid", fgColor="FEF3C7")
    garis = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    for sel in ws[1]:
        sel.fill = isi
        sel.font = Font(bold=True, color="FFFFFF")
        sel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for nomor in range(1, 41):
        ws.append([nomor, None, None, None, None])
    for baris in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(judul)):
        for sel in baris:
            sel.border = garis
            if sel.row == 1:
                continue
            sel.alignment = Alignment(vertical="center")
            if sel.row <= 1 + len(contoh):
                sel.fill = kuning
                sel.font = Font(italic=True, color="92400E")
    lebar = (12, 36, 28, 14, 36)
    for indeks, ukuran in enumerate(lebar, start=1):
        ws.column_dimensions[get_column_letter(indeks)].width = ukuran
    ws.auto_filter.ref = f"A1:E{ws.max_row}"
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = "Form Data Guru"
    ws.oddFooter.center.text = "SMP Muhammadiyah 22 Setiabudi Pamulang"
    keluar = io.BytesIO()
    wb.save(keluar)
    keluar.seek(0)
    return keluar


def parse_guru_file(filename, data):
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        sheets = _rows_from_xlsx(data)
    elif name.endswith(".xls"):
        sheets = _rows_from_xls(data)
    else:
        raise ValueError("File harus berformat .xls atau .xlsx seperti jadwal pelajaran.")
    return parse_guru_rows(sheets)


def ensure_guru_columns():
    cols = {row["Field"] for row in query("SHOW COLUMNS FROM guru")}
    if "mapel_ajar" not in cols:
        execute("ALTER TABLE guru ADD COLUMN mapel_ajar TEXT NULL")
    if "jumlah_jam" not in cols:
        execute("ALTER TABLE guru ADD COLUMN jumlah_jam INT NULL")
    if "mengajar_di" not in cols:
        execute("ALTER TABLE guru ADD COLUMN mengajar_di TEXT NULL")
    if "nbm" not in cols:
        execute("ALTER TABLE guru ADD COLUMN nbm VARCHAR(30) NULL")
    execute("ALTER TABLE guru MODIFY jenis_kelamin ENUM('L','P') NULL")
    execute("ALTER TABLE guru MODIFY gelar_belakang VARCHAR(80) DEFAULT ''")


def replace_guru(records):
    ensure_guru_columns()
    existing = query("SELECT id, nama FROM guru")
    by_name = {re.sub(r"\s+", " ", row["nama"]).strip().lower(): row["id"] for row in existing}
    kept = []
    with transaction():
        for rec in records:
            key = rec["nama"].strip().lower()
            if key in by_name:
                guru_id = by_name[key]
                execute(
                    """UPDATE guru
                       SET nama=%s, gelar_belakang=%s, mapel_ajar=%s, jumlah_jam=%s, mengajar_di=%s, aktif=1
                       WHERE id=%s""",
                    (rec["nama"], rec["gelar_belakang"], rec["mapel_ajar"], rec["jumlah_jam"], rec["mengajar_di"], guru_id),
                )
            else:
                guru_id = execute(
                    """INSERT INTO guru
                       (nama, gelar_belakang, mapel_ajar, jumlah_jam, mengajar_di, jenis_kelamin, status_kepegawaian, aktif)
                       VALUES (%s,%s,%s,%s,%s,NULL,'Guru',1)""",
                    (rec["nama"], rec["gelar_belakang"], rec["mapel_ajar"], rec["jumlah_jam"], rec["mengajar_di"]),
                )
                by_name[key] = guru_id
            kept.append(guru_id)
        placeholders = ", ".join(["%s"] * len(kept))
        execute(f"DELETE FROM guru WHERE id NOT IN ({placeholders})", kept)
    return len(kept)
