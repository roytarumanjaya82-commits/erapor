import os
from io import BytesIO
from zipfile import ZipFile

from flask import Blueprint, flash, redirect, render_template, request, send_file, session, url_for
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from werkzeug.security import generate_password_hash

from .access import ENDPOINT_TUGAS
from .auth import login_required
from .db import execute, query, transaction

bp = Blueprint("menu", __name__)

ENDPOINT_TUGAS.update(
    {
        "menu.audit": {"admin", "bidang_studi"},
        "menu.leger": {"admin", "walas", "bidang_studi"},
        "menu.leger_unduh": {"admin", "walas", "bidang_studi"},
        "menu.leger_unduh_semua": {"admin", "walas", "bidang_studi"},
        "menu.cetak_rapor": {"admin", "walas"},
        "menu.rapor_siswa": {"admin", "walas"},
        "menu.publik": {"admin", "bidang_studi", "walas", "ekskul", "tahfidz"},
        "menu.dapodik": {"admin"},
        "menu.pengguna": {"admin"},
        "menu.form_pengguna": {"admin"},
        "menu.hapus_pengguna": {"admin"},
        "menu.generate_semua": {"admin"},
        "menu.generate_guru": {"admin"},
        "menu.generate_siswa": {"admin"},
        "menu.kosongkan": {"admin"},
        "menu.backup": {"admin"},
        "menu.backup_buat": {"admin"},
        "menu.backup_file": {"admin"},
        "menu.backup_hapus": {"admin"},
        "menu.backup_unduh": {"admin"},
        "menu.restore": {"admin"},
        "menu.dimensi": {"admin", "walas"},
        "menu.simpan_dimensi": {"admin"},
        "menu.toggle_dimensi": {"admin"},
        "menu.hapus_dimensi": {"admin"},
        "menu.predikat": {"admin", "walas"},
        "menu.simpan_predikat": {"admin"},
        "menu.toggle_predikat": {"admin"},
        "menu.hapus_predikat": {"admin"},
        "menu.cetak_tahfiz": {"admin", "tahfidz", "walas"},
        "menu.nilai_mapel": {"admin", "bidang_studi"},
        "menu.simpan_nilai_mapel": {"admin", "bidang_studi"},
        "menu.nilai_ekskul": {"admin", "ekskul"},
        "menu.simpan_ekskul": {"admin", "ekskul"},
    }
)


def _ta_id():
    ta = query("SELECT id FROM tahun_ajaran WHERE aktif=1 LIMIT 1", one=True)
    return ta["id"] if ta else None


def _is_admin():
    return "admin" in set(session.get("tugas") or [])


def _guru_id():
    return session.get("guru_id")


def _kelas_scope():
    if _is_admin():
        return query(
            "SELECT id, nama FROM kelas ORDER BY FIELD(tingkat,'7','8','9'), nama"
        )
    gid = _guru_id()
    if not gid:
        return []
    tugas = set(session.get("tugas") or [])
    rows = []
    if "walas" in tugas:
        rows += query("SELECT id, nama FROM kelas WHERE wali_kelas_id=%s", (gid,))
    if "tahfidz" in tugas:
        rows += query(
            """SELECT DISTINCT k.id, k.nama FROM kelas k
               JOIN mapping_guru_tahfiz m ON m.kelas_id=k.id
               WHERE m.guru_id=%s""",
            (gid,),
        )
    if "bidang_studi" in tugas:
        rows += query(
            """SELECT DISTINCT k.id, k.nama FROM kelas k
               JOIN pelajaran p ON p.kelas_id=k.id
               WHERE p.guru_id=%s""",
            (gid,),
        )
    unik = {}
    for row in rows:
        unik[row["id"]] = row
    return list(unik.values())


@bp.route("/monitoring/publik")
@login_required
def publik():
    ta_id = _ta_id()
    pelajaran = query("SELECT COUNT(*) AS n FROM pelajaran WHERE tahun_ajaran_id=%s", (ta_id,), one=True)["n"] if ta_id else 0
    lengkap = query(
        """SELECT COUNT(*) AS n FROM pelajaran pel
           WHERE pel.tahun_ajaran_id=%s
             AND (SELECT COUNT(*) FROM siswa s WHERE s.kelas_id=pel.kelas_id AND s.status='aktif') > 0
             AND (SELECT COUNT(*) FROM nilai_sumatif ns WHERE ns.pelajaran_id=pel.id)
                 >= (SELECT COUNT(*) FROM siswa s WHERE s.kelas_id=pel.kelas_id AND s.status='aktif')""",
        (ta_id,),
        one=True,
    )["n"] if ta_id else 0
    stats = {
        "pelajaran": pelajaran,
        "lengkap": lengkap,
        "pct": round(lengkap * 100 / pelajaran) if pelajaran else 0,
        "guru": query("SELECT COUNT(*) AS n FROM guru WHERE aktif=1", one=True)["n"],
        "kelas": query("SELECT COUNT(*) AS n FROM kelas", one=True)["n"],
        "siswa": query("SELECT COUNT(*) AS n FROM siswa WHERE status='aktif'", one=True)["n"],
    }
    return render_template("menu/publik.html", stats=stats, nav="publik")


@bp.route("/admin/audit")
@login_required
def audit():
    ta_id = _ta_id()
    sql = """
        SELECT k.nama AS kelas_nama, m.nama AS mapel_nama, g.nama AS guru_nama,
               (SELECT COUNT(*) FROM siswa s WHERE s.kelas_id=pel.kelas_id AND s.status='aktif') AS siswa,
               (SELECT COUNT(*) FROM nilai_formatif nf WHERE nf.pelajaran_id=pel.id) AS formatif,
               (SELECT COUNT(*) FROM nilai_sumatif ns WHERE ns.pelajaran_id=pel.id) AS sumatif
        FROM pelajaran pel
        JOIN kelas k ON k.id=pel.kelas_id
        JOIN mata_pelajaran m ON m.id=pel.mapel_id
        LEFT JOIN guru g ON g.id=pel.guru_id
        WHERE pel.tahun_ajaran_id=%s
    """
    params = [ta_id]
    if not _is_admin():
        sql += " AND pel.guru_id=%s"
        params.append(_guru_id() or 0)
    sql += " ORDER BY k.nama, m.nama"
    rows = query(sql, params) if ta_id else []
    for row in rows:
        total = row["siswa"] or 0
        row["fmt"] = round(row["formatif"] * 100 / total) if total else 0
        row["sum"] = round(row["sumatif"] * 100 / total) if total else 0
        if total and row["formatif"] >= total and row["sumatif"] >= total:
            row["status"] = "Lengkap"
        elif row["formatif"] or row["sumatif"]:
            row["status"] = "Proses"
        else:
            row["status"] = "Belum mengisi"
    return render_template("menu/audit.html", rows=rows, nav="audit")


def _angka(nilai):
    if nilai is None:
        return None
    return float(nilai)


def _teks_nilai(nilai):
    if nilai is None:
        return ""
    return format(nilai, "f").rstrip("0").rstrip(".")


def _susun_leger(kelas_id):
    ta_id = _ta_id()
    info = query(
        """SELECT k.id, k.nama, k.tingkat, g.nama AS wali, g.gelar_belakang,
                  (SELECT COUNT(*) FROM siswa s WHERE s.kelas_id=k.id AND s.status='aktif') AS siswa
           FROM kelas k
           LEFT JOIN guru g ON g.id=k.wali_kelas_id
           WHERE k.id=%s""",
        (kelas_id,),
        one=True,
    )
    if not info:
        return None
    if info["gelar_belakang"] and info["wali"] and info["gelar_belakang"] not in info["wali"]:
        info["wali"] = f"{info['wali']}, {info['gelar_belakang']}"
    mapel = query(
        """SELECT pel.id, m.nama, COALESCE(pp.kkm, m.kkm_default, 75) AS kkm
           FROM pelajaran pel
           JOIN mata_pelajaran m ON m.id=pel.mapel_id
           LEFT JOIN perencanaan_penilaian pp ON pp.pelajaran_id=pel.id
           LEFT JOIN mapping_rapor mr ON mr.id = (
             SELECT x.id FROM mapping_rapor x
             WHERE x.mapel_id=m.id AND x.tingkat=%s
             ORDER BY x.id LIMIT 1
           )
           WHERE pel.kelas_id=%s AND (%s IS NULL OR pel.tahun_ajaran_id=%s)
           ORDER BY COALESCE(mr.urutan, 999), m.nama""",
        (info["tingkat"], kelas_id, ta_id, ta_id),
    ) if ta_id else []
    siswa = query(
        "SELECT id, nis, nisn, nama, jenis_kelamin FROM siswa WHERE kelas_id=%s AND status='aktif' ORDER BY nama",
        (kelas_id,),
    )
    formatif = {
        (row["siswa_id"], row["pelajaran_id"]): row["nilai"]
        for row in query(
            """SELECT nf.siswa_id, nf.pelajaran_id, nf.nilai FROM nilai_formatif nf
               JOIN pelajaran pel ON pel.id=nf.pelajaran_id WHERE pel.kelas_id=%s""",
            (kelas_id,),
        )
    }
    sumatif = {
        (row["siswa_id"], row["pelajaran_id"]): row["nilai"]
        for row in query(
            """SELECT ns.siswa_id, ns.pelajaran_id, ns.nilai FROM nilai_sumatif ns
               JOIN pelajaran pel ON pel.id=ns.pelajaran_id WHERE pel.kelas_id=%s""",
            (kelas_id,),
        )
    }
    for item in siswa:
        ringkas = []
        rinci = []
        angka = []
        for pelajaran in mapel:
            kunci = (item["id"], pelajaran["id"])
            f_nilai = _angka(formatif.get(kunci))
            s_nilai = _angka(sumatif.get(kunci))
            tersedia = [n for n in (f_nilai, s_nilai) if n is not None]
            akhir = round(sum(tersedia) / len(tersedia), 1) if tersedia else None
            if akhir is not None:
                angka.append(akhir)
            kkm = float(pelajaran["kkm"] or 75)
            ringkas.append({"teks": _teks_nilai(akhir) if akhir is not None else "", "rendah": akhir is not None and akhir < kkm})
            rinci.append(
                {
                    "formatif": _teks_nilai(f_nilai) if f_nilai is not None else "",
                    "sumatif": _teks_nilai(s_nilai) if s_nilai is not None else "",
                    "na": _teks_nilai(akhir) if akhir is not None else "",
                    "rendah": akhir is not None and akhir < kkm,
                }
            )
        item["nilai"] = ringkas
        item["rinci"] = rinci
        item["jumlah"] = round(sum(angka), 1) if angka else None
        item["rata"] = round(sum(angka) / len(angka), 1) if angka else None
        item["peringkat"] = None
    berperingkat = sorted((item for item in siswa if item["rata"] is not None), key=lambda item: (-item["rata"], item["nama"]))
    urutan = 0
    sebelumnya = None
    for indeks, item in enumerate(berperingkat, start=1):
        if item["rata"] != sebelumnya:
            urutan = indeks
            sebelumnya = item["rata"]
        item["peringkat"] = urutan
    rata_siswa = [item["rata"] for item in siswa if item["rata"] is not None]
    info["mapel"] = len(mapel)
    info["rata"] = round(sum(rata_siswa) / len(rata_siswa), 1) if rata_siswa else None
    return {"info": info, "mapel": mapel, "siswa": siswa}


def _kelas_terpilih():
    kelas = _kelas_scope()
    boleh = {str(item["id"]) for item in kelas}
    minta = request.args.get("kelas_id") or ""
    kelas_id = minta if minta in boleh else (kelas[0]["id"] if kelas else "")
    return kelas, str(kelas_id) if kelas_id else ""


def _nama_berkas(nama):
    aman = "".join(ch if ch.isalnum() or ch in "-_" else "-" for ch in nama)
    return aman.strip("-") or "kelas"


def _isi_excel(paket, mode="semua"):
    wb = Workbook()
    tipis = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB"),
    )
    kepala = PatternFill("solid", fgColor="0F2744")
    huruf = Font(bold=True, color="FFFFFF")
    info = paket["info"]

    def tulis(ws, lengkap):
        ws["A1"] = f"Leger Nilai {info['nama']}"
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([])
        if lengkap:
            atas = ["No", "NISN", "Nama Siswa", "L/P"]
            bawah = ["", "", "", ""]
            for pelajaran in paket["mapel"]:
                atas.extend([pelajaran["nama"], "", ""])
                bawah.extend(["Formatif", "Sumatif", "NA"])
            atas.extend(["Jumlah", "Rata-rata", "Peringkat"])
            bawah.extend(["", "", ""])
            ws.append(atas)
            ws.append(bawah)
            kolom = 5
            for _pelajaran in paket["mapel"]:
                ws.merge_cells(start_row=3, start_column=kolom, end_row=3, end_column=kolom + 2)
                kolom += 3
            baris_kepala = 4
        else:
            kepala_baris = ["No", "NISN", "Nama Siswa", "L/P"]
            kepala_baris.extend(pelajaran["nama"] for pelajaran in paket["mapel"])
            kepala_baris.extend(["Jumlah", "Rata-rata", "Peringkat"])
            ws.append(kepala_baris)
            baris_kepala = 3
        for sel in ws[baris_kepala]:
            sel.fill = kepala
            sel.font = huruf
            sel.alignment = Alignment(horizontal="center", wrap_text=True)
        if lengkap:
            for sel in ws[3]:
                sel.fill = kepala
                sel.font = huruf
                sel.alignment = Alignment(horizontal="center")
        for nomor, siswa in enumerate(paket["siswa"], start=1):
            baris = [nomor, siswa["nisn"] or siswa["nis"], siswa["nama"], siswa["jenis_kelamin"] or ""]
            if lengkap:
                for nilai in siswa["rinci"]:
                    baris.extend([nilai["formatif"] or None, nilai["sumatif"] or None, nilai["na"] or None])
            else:
                baris.extend(nilai["teks"] or None for nilai in siswa["nilai"])
            baris.extend([siswa["jumlah"], siswa["rata"], siswa["peringkat"]])
            ws.append(baris)
        for baris in ws.iter_rows(min_row=baris_kepala, max_row=ws.max_row, max_col=ws.max_column):
            for sel in baris:
                sel.border = tipis
        ws.freeze_panes = "E5" if lengkap else "E4"
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["C"].width = 32

    if mode == "lengkap":
        lengkap = wb.active
        lengkap.title = "Rincian Lengkap"
        tulis(lengkap, True)
    else:
        ringkas = wb.active
        ringkas.title = "Nilai Rapor"
        tulis(ringkas, False)
        if mode != "ringkas":
            tulis(wb.create_sheet("Rincian Lengkap"), True)
    berkas = BytesIO()
    wb.save(berkas)
    berkas.seek(0)
    return berkas


@bp.route("/admin/leger")
@login_required
def leger():
    kelas, kelas_id = _kelas_terpilih()
    tampil = request.args.get("tampil") or "pratinjau"
    if tampil not in {"pratinjau", "daftar"}:
        tampil = "pratinjau"
    mode = request.args.get("mode") or "ringkas"
    if mode not in {"ringkas", "lengkap"}:
        mode = "ringkas"
    paket = _susun_leger(kelas_id) if kelas_id and tampil == "pratinjau" else None
    daftar = []
    if tampil == "daftar":
        ta_id = _ta_id()
        for item in kelas:
            ringkas = query(
                """SELECT k.nama, k.tingkat, g.nama AS wali,
                          (SELECT COUNT(*) FROM siswa s WHERE s.kelas_id=k.id AND s.status='aktif') AS siswa,
                          (SELECT COUNT(*) FROM pelajaran p WHERE p.kelas_id=k.id AND (%s IS NULL OR p.tahun_ajaran_id=%s)) AS mapel
                   FROM kelas k LEFT JOIN guru g ON g.id=k.wali_kelas_id WHERE k.id=%s""",
                (ta_id, ta_id, item["id"]),
                one=True,
            )
            if ringkas:
                daftar.append(ringkas | {"id": item["id"]})
    return render_template(
        "menu/leger.html",
        kelas=kelas,
        kelas_id=kelas_id,
        tampil=tampil,
        mode=mode,
        paket=paket,
        daftar=daftar,
        nav="leger",
    )


@bp.route("/admin/leger/unduh")
@login_required
def leger_unduh():
    kelas, kelas_id = _kelas_terpilih()
    minta = request.args.get("kelas_id") or ""
    if minta and minta not in {str(item["id"]) for item in kelas}:
        flash("Kelas ini tidak termasuk tugas Anda.", "error")
        return redirect(url_for("menu.leger"))
    if not kelas_id:
        flash("Belum ada kelas pada tugas Anda.", "error")
        return redirect(url_for("menu.leger"))
    paket = _susun_leger(kelas_id)
    if not paket:
        flash("Kelas tidak ditemukan.", "error")
        return redirect(url_for("menu.leger"))
    mode = request.args.get("mode") or "semua"
    if mode not in {"semua", "ringkas", "lengkap"}:
        mode = "semua"
    return send_file(
        _isi_excel(paket, mode),
        as_attachment=True,
        download_name=f"leger-{_nama_berkas(paket['info']['nama'])}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@bp.route("/admin/leger/unduh-semua")
@login_required
def leger_unduh_semua():
    kelas = _kelas_scope()
    if not kelas:
        flash("Belum ada kelas pada tugas Anda.", "error")
        return redirect(url_for("menu.leger"))
    arsip = BytesIO()
    with ZipFile(arsip, "w") as zip_file:
        for item in kelas:
            paket = _susun_leger(item["id"])
            if not paket:
                continue
            zip_file.writestr(f"leger-{_nama_berkas(paket['info']['nama'])}.xlsx", _isi_excel(paket).getvalue())
    arsip.seek(0)
    return send_file(arsip, as_attachment=True, download_name="leger-semua-kelas.zip", mimetype="application/zip")


@bp.route("/admin/cetak-rapor")
@login_required
def cetak_rapor():
    if _is_admin():
        kelas = _kelas_scope()
    else:
        kelas = query(
            "SELECT id, nama FROM kelas WHERE wali_kelas_id=%s ORDER BY nama",
            (_guru_id() or 0,),
        )
    boleh = {str(item["id"]) for item in kelas}
    kelas_id = request.args.get("kelas_id") or ""
    if str(kelas_id) not in boleh:
        kelas_id = ""
    siswa = query(
        "SELECT id, nis, nisn, nama, jenis_kelamin FROM siswa WHERE kelas_id=%s AND status='aktif' ORDER BY nama",
        (kelas_id,),
    ) if kelas_id else []
    return render_template("menu/cetak_rapor.html", kelas=kelas, kelas_id=str(kelas_id), siswa=siswa, nav="cetak")


@bp.route("/admin/cetak-rapor/<int:siswa_id>")
@login_required
def rapor_siswa(siswa_id):
    siswa = query(
        """SELECT s.*, k.nama AS kelas_nama, g.nama AS wali, g.gelar_belakang
           FROM siswa s LEFT JOIN kelas k ON k.id=s.kelas_id
           LEFT JOIN guru g ON g.id=k.wali_kelas_id WHERE s.id=%s""",
        (siswa_id,),
        one=True,
    )
    if not siswa:
        flash("Siswa tidak ditemukan.", "error")
        return redirect(url_for("menu.cetak_rapor"))
    if not _is_admin():
        milik = query(
            "SELECT id FROM kelas WHERE id=%s AND wali_kelas_id=%s",
            (siswa["kelas_id"], _guru_id() or 0),
            one=True,
        )
        if not milik:
            flash("Anda hanya dapat mencetak rapor kelas binaannya.", "error")
            return redirect(url_for("menu.cetak_rapor"))
    mapel = query(
        """SELECT m.nama, p.kkm, ns.nilai
           FROM pelajaran pel
           JOIN mata_pelajaran m ON m.id=pel.mapel_id
           LEFT JOIN perencanaan_penilaian p ON p.pelajaran_id=pel.id
           LEFT JOIN nilai_sumatif ns ON ns.pelajaran_id=pel.id AND ns.siswa_id=%s
           WHERE pel.kelas_id=%s ORDER BY m.nama""",
        (siswa_id, siswa["kelas_id"]),
    )
    ekskul = query(
        """SELECT e.nama, n.nilai FROM nilai_ekskul n
           JOIN ekstrakurikuler e ON e.id=n.ekskul_id WHERE n.siswa_id=%s""",
        (siswa_id,),
    )
    sekolah = query("SELECT * FROM sekolah LIMIT 1", one=True) or {}
    ta = query("SELECT * FROM tahun_ajaran WHERE aktif=1 LIMIT 1", one=True)
    return render_template("menu/rapor_siswa.html", siswa=siswa, mapel=mapel, ekskul=ekskul, sekolah=sekolah, ta=ta)


@bp.route("/admin/dapodik")
@login_required
def dapodik():
    return render_template("menu/dapodik.html", nav="dapodik")


def _email_pengguna(username, email):
    bersih = (email or "").strip()
    return bersih or f"{username}@erapor.local"


def _simpan_tugas(user_id, role, tugas):
    execute("DELETE FROM user_tugas WHERE user_id=%s", (user_id,))
    if role != "guru":
        return
    for item in tugas:
        if item in {"bidang_studi", "walas", "ekskul", "tahfidz"}:
            execute("INSERT INTO user_tugas (user_id, tugas) VALUES (%s,%s)", (user_id, item))


@bp.route("/admin/pengguna")
@login_required
def pengguna():
    from .access import ensure_access

    ensure_access()
    tab = request.args.get("tab") or "admin"
    if tab not in {"admin", "guru", "siswa"}:
        tab = "admin"
    q = (request.args.get("q") or "").strip()
    try:
        per = int(request.args.get("per") or 10)
    except ValueError:
        per = 10
    if per not in {10, 25, 50, 100}:
        per = 10
    try:
        halaman = max(int(request.args.get("halaman") or 1), 1)
    except ValueError:
        halaman = 1
    where = ["u.role=%s"]
    params = [tab if tab != "guru" else "guru"]
    if tab == "guru":
        where = ["u.role NOT IN ('admin','siswa')"]
        params = []
    if q:
        where.append("(u.username LIKE %s OR u.nama LIKE %s OR u.email LIKE %s)")
        like = f"%{q}%"
        params.extend([like, like, like])
    syarat = " AND ".join(where)
    total = query(f"SELECT COUNT(*) AS n FROM users u WHERE {syarat}", params, one=True)["n"]
    halaman_max = max((total + per - 1) // per, 1)
    if halaman > halaman_max:
        halaman = halaman_max
    offset = (halaman - 1) * per
    rows = query(
        f"""SELECT u.*, GROUP_CONCAT(t.tugas ORDER BY t.tugas SEPARATOR ', ') AS tugas_nama
            FROM users u LEFT JOIN user_tugas t ON t.user_id=u.id
            WHERE {syarat}
            GROUP BY u.id ORDER BY u.nama LIMIT %s OFFSET %s""",
        (*params, per, offset),
    )
    return render_template(
        "menu/pengguna.html",
        rows=rows,
        tab=tab,
        q=q,
        per=per,
        halaman=halaman,
        total=total,
        awal=0 if total == 0 else offset + 1,
        akhir=min(offset + per, total),
        nav="pengguna",
    )


@bp.route("/admin/pengguna/form", methods=["GET", "POST"])
@bp.route("/admin/pengguna/form/<int:item_id>", methods=["GET", "POST"])
@login_required
def form_pengguna(item_id=None):
    from .access import ensure_access

    ensure_access()
    row = query("SELECT * FROM users WHERE id=%s", (item_id,), one=True) if item_id else None
    if item_id and not row:
        flash("Pengguna tidak ditemukan.", "error")
        return redirect(url_for("menu.pengguna"))
    tugas_aktif = set()
    if row:
        tugas_aktif = {item["tugas"] for item in query("SELECT tugas FROM user_tugas WHERE user_id=%s", (row["id"],))}
    if request.method == "POST":
        username = (request.form.get("username") or "").strip().lower()
        nama = (request.form.get("nama") or "").strip()
        role = request.form.get("role") or "guru"
        if role not in {"admin", "guru", "siswa"}:
            role = "guru"
        password = request.form.get("password") or ""
        guru_id = request.form.get("guru_id") or None
        aktif = 1 if request.form.get("aktif") != "0" else 0
        if not username or not nama:
            flash("Username dan nama wajib diisi.", "error")
            return redirect(request.url)
        if row and row["username"] == "admin" and (role != "admin" or aktif == 0):
            flash("Akun admin utama harus tetap admin dan aktif.", "error")
            return redirect(request.url)
        bentrok = query(
            "SELECT id FROM users WHERE username=%s AND id<>%s",
            (username, row["id"] if row else 0),
            one=True,
        )
        if bentrok:
            flash("Username sudah dipakai.", "error")
            return redirect(request.url)
        email = _email_pengguna(username, request.form.get("email"))
        if row:
            if password:
                execute(
                    """UPDATE users SET username=%s,nama=%s,email=%s,role=%s,guru_id=%s,aktif=%s,password_hash=%s
                       WHERE id=%s""",
                    (username, nama, email, role, guru_id, aktif, generate_password_hash(password), row["id"]),
                )
            else:
                execute(
                    "UPDATE users SET username=%s,nama=%s,email=%s,role=%s,guru_id=%s,aktif=%s WHERE id=%s",
                    (username, nama, email, role, guru_id, aktif, row["id"]),
                )
            _simpan_tugas(row["id"], role, request.form.getlist("tugas"))
            flash("Pengguna diperbarui.", "success")
        else:
            uid = execute(
                """INSERT INTO users (username, password_hash, nama, email, role, guru_id, aktif)
                   VALUES (%s,%s,%s,%s,%s,%s,%s)""",
                (username, generate_password_hash(password or "guru123"), nama, email, role, guru_id, aktif),
            )
            _simpan_tugas(uid, role, request.form.getlist("tugas"))
            flash("Pengguna ditambahkan.", "success")
        return redirect(url_for("menu.pengguna", tab="guru" if role == "guru" else role))
    guru = query("SELECT id, nama FROM guru WHERE aktif=1 ORDER BY nama")
    return render_template(
        "menu/pengguna_form.html",
        row=row or {},
        tugas_aktif=tugas_aktif,
        guru=guru,
        nav="pengguna",
    )


@bp.route("/admin/pengguna/hapus/<int:item_id>", methods=["POST"])
@login_required
def hapus_pengguna(item_id):
    row = query("SELECT id, username, role FROM users WHERE id=%s", (item_id,), one=True)
    if not row:
        flash("Pengguna tidak ditemukan.", "error")
        return redirect(url_for("menu.pengguna"))
    if row["username"] == "admin" or row["id"] == session.get("user_id"):
        flash("Akun ini tidak dapat dihapus.", "error")
        return redirect(url_for("menu.pengguna", tab=row["role"] if row["role"] in {"admin", "siswa"} else "guru"))
    execute("DELETE FROM users WHERE id=%s", (item_id,))
    flash("Pengguna dihapus.", "success")
    return redirect(url_for("menu.pengguna", tab=row["role"] if row["role"] in {"admin", "siswa"} else "guru"))


@bp.route("/admin/pengguna/generate-semua", methods=["POST"])
@login_required
def generate_semua():
    guru = _buat_akun_guru()
    siswa = _buat_akun_siswa()
    flash(f"{guru} akun guru dan {siswa} akun siswa baru dibuat. Kata sandi guru: guru123. Kata sandi siswa sama dengan NIS.", "success")
    return redirect(url_for("menu.pengguna", tab=request.args.get("tab") or "admin"))


def _buat_akun_guru():
    dibuat = 0
    for guru in query("SELECT id, nama FROM guru"):
        username = "".join(ch for ch in guru["nama"].lower() if ch.isalnum())[:20] or f"guru{guru['id']}"
        if query("SELECT id FROM users WHERE username=%s OR guru_id=%s", (username, guru["id"]), one=True):
            continue
        execute(
            """INSERT INTO users (username, password_hash, nama, email, role, guru_id, aktif)
               VALUES (%s,%s,%s,%s,'guru',%s,1)""",
            (username, generate_password_hash("guru123"), guru["nama"], f"{username}@erapor.local", guru["id"]),
        )
        dibuat += 1
    return dibuat


def _buat_akun_siswa():
    dibuat = 0
    for siswa in query("SELECT id, nis, nama FROM siswa WHERE status='aktif'"):
        username = f"siswa{siswa['nis']}"[:50]
        if query("SELECT id FROM users WHERE username=%s", (username,), one=True):
            continue
        execute(
            """INSERT INTO users (username, password_hash, nama, email, role, aktif)
               VALUES (%s,%s,%s,%s,'siswa',1)""",
            (username, generate_password_hash(str(siswa["nis"])), siswa["nama"], f"{username}@erapor.local"),
        )
        dibuat += 1
    return dibuat


@bp.route("/admin/pengguna/generate-guru", methods=["POST"])
@login_required
def generate_guru():
    dibuat = _buat_akun_guru()
    flash(f"{dibuat} akun guru dibuat. Kata sandi awal: guru123. Atur tugas tiap guru setelah itu.", "success")
    return redirect(url_for("menu.pengguna", tab="guru"))


@bp.route("/admin/pengguna/generate-siswa", methods=["POST"])
@login_required
def generate_siswa():
    dibuat = _buat_akun_siswa()
    flash(f"{dibuat} akun siswa dibuat. Username siswa+NIS, kata sandi sama dengan NIS.", "success")
    return redirect(url_for("menu.pengguna", tab="siswa"))


def _tabel_ada(nama):
    return query(
        "SELECT COUNT(*) AS n FROM information_schema.TABLES WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME=%s",
        (nama,),
        one=True,
    )["n"]


def _jumlah_tabel(nama):
    if not _tabel_ada(nama):
        return 0
    return query(f"SELECT COUNT(*) AS n FROM `{nama}`", one=True)["n"]


def _hapus_tabel(daftar):
    for nama in daftar:
        if _tabel_ada(nama):
            execute(f"DELETE FROM `{nama}`")


def _hapus_akun_non_admin():
    execute("DELETE FROM users WHERE role<>'admin' AND username<>'admin'")


AKSI_KOSONGKAN = {
    "nilai": (
        "Data nilai dikosongkan.",
        ["nilai_formatif", "nilai_sumatif", "nilai_tahfiz", "nilai_ekskul", "nilai_tryout", "perencanaan_penilaian"],
    ),
    "siswa": ("Data siswa dikosongkan.", ["siswa"]),
    "guru": ("Data guru dan akun non-admin dikosongkan. Akun administrator tetap ada.", None),
    "rombel": ("Data rombel, ekskul, dan pembelajaran dikosongkan.", ["ekstrakurikuler", "kelas"]),
    "reset": ("Seluruh data aplikasi dikosongkan. Akun administrator tetap ada.", None),
}


@bp.route("/admin/kosongkan", methods=["GET", "POST"])
@login_required
def kosongkan():
    if request.method == "POST":
        aksi = request.form.get("aksi")
        if aksi not in AKSI_KOSONGKAN:
            flash("Kategori pengosongan tidak dikenal.", "error")
            return redirect(url_for("menu.kosongkan"))
        pesan, tabel = AKSI_KOSONGKAN[aksi]
        with transaction():
            if aksi == "guru":
                _hapus_akun_non_admin()
                _hapus_tabel(["guru"])
            elif aksi == "reset":
                _hapus_tabel(
                    [
                        "nilai_formatif", "nilai_sumatif", "nilai_tahfiz", "nilai_ekskul", "nilai_tryout",
                        "perencanaan_penilaian", "pelajaran", "mapping_guru_tahfiz", "siswa", "kelas",
                        "ekstrakurikuler", "gabung_mapel_item", "gabung_mapel", "mapping_rapor",
                        "mata_pelajaran", "kelompok_mapel", "tryout", "aspek_tahfiz",
                    ]
                )
                _hapus_akun_non_admin()
                _hapus_tabel(["guru"])
            else:
                _hapus_tabel(tabel)
        flash(pesan, "success")
        return redirect(url_for("menu.kosongkan"))
    jumlah = {
        "nilai": sum(
            _jumlah_tabel(nama)
            for nama in ("nilai_formatif", "nilai_sumatif", "nilai_tahfiz", "nilai_ekskul", "perencanaan_penilaian")
        ),
        "siswa": _jumlah_tabel("siswa"),
        "guru": _jumlah_tabel("guru"),
        "rombel": _jumlah_tabel("kelas"),
    }
    return render_template("menu/kosongkan.html", jumlah=jumlah, nav="kosongkan")


def _sql_backup():
    tables = [
        "sekolah", "tahun_ajaran", "guru", "users", "user_tugas", "kelas", "siswa",
        "kelompok_mapel", "mata_pelajaran", "pelajaran", "perencanaan_penilaian",
        "nilai_formatif", "nilai_sumatif", "ekstrakurikuler", "nilai_ekskul",
        "gabung_mapel", "gabung_mapel_item", "mapping_rapor", "tryout",
        "aspek_tahfiz", "mapping_guru_tahfiz", "pengaturan_rapor_tahfiz", "nilai_tahfiz",
        "dimensi_sikap", "predikat_sikap",
    ]
    chunks = ["-- eRapor backup SMPM22SETIABUDIPAMULANG\nSET FOREIGN_KEY_CHECKS=0;\n"]
    for table in tables:
        try:
            rows = query(f"SELECT * FROM {table}")
        except Exception:
            continue
        chunks.append(f"DELETE FROM {table};\n")
        if not rows:
            continue
        cols = list(rows[0].keys())
        col_sql = ", ".join(f"`{col}`" for col in cols)
        for row in rows:
            vals = []
            for col in cols:
                value = row[col]
                if value is None:
                    vals.append("NULL")
                elif isinstance(value, (int, float)):
                    vals.append(str(value))
                else:
                    text = str(value).replace("\\", "\\\\").replace("'", "''")
                    vals.append(f"'{text}'")
            chunks.append(f"INSERT INTO {table} ({col_sql}) VALUES ({', '.join(vals)});\n")
    chunks.append("SET FOREIGN_KEY_CHECKS=1;\n")
    return "".join(chunks).encode("utf-8")


def _pecah_sql(text):
    bagian = []
    buf = []
    kutip = False
    i = 0
    while i < len(text):
        ch = text[i]
        if ch == "'" and kutip and i + 1 < len(text) and text[i + 1] == "'":
            buf.append("''")
            i += 2
            continue
        if ch == "'":
            kutip = not kutip
        if ch == ";" and not kutip:
            bagian.append("".join(buf))
            buf = []
        else:
            buf.append(ch)
        i += 1
    if "".join(buf).strip():
        bagian.append("".join(buf))
    return bagian


def _jalankan_sql(text):
    from .db import get_conn

    conn = get_conn()
    conn.autocommit(False)
    try:
        with conn.cursor() as cur:
            for statement in _pecah_sql(text):
                lines = statement.splitlines()
                while lines and (not lines[0].strip() or lines[0].strip().startswith("--")):
                    lines.pop(0)
                sql = "\n".join(lines).strip()
                if sql:
                    cur.execute(sql)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.autocommit(True)


@bp.route("/admin/backup", methods=["GET"])
@login_required
def backup():
    from .arsip import daftar_backup

    return render_template("menu/backup.html", berkas=daftar_backup(), nav="backup")


@bp.route("/admin/backup/buat", methods=["POST"])
@login_required
def backup_buat():
    from .arsip import bungkus, simpan_backup

    blob = bungkus(_sql_backup(), request.form.get("kunci"))
    nama = simpan_backup(blob)
    return send_file(
        BytesIO(blob),
        as_attachment=True,
        download_name=nama,
        mimetype="application/octet-stream",
    )


@bp.route("/admin/backup/unduh")
@login_required
def backup_unduh():
    from .arsip import bungkus, simpan_backup

    blob = bungkus(_sql_backup(), "")
    nama = simpan_backup(blob)
    return send_file(BytesIO(blob), as_attachment=True, download_name=nama, mimetype="application/octet-stream")


@bp.route("/admin/backup/file/<nama>")
@login_required
def backup_file(nama):
    from .arsip import folder_backup, nama_aman

    aman = nama_aman(nama)
    jalur = os.path.join(folder_backup(), aman) if aman else ""
    if not aman or not os.path.isfile(jalur):
        flash("File backup tidak ditemukan.", "error")
        return redirect(url_for("menu.backup"))
    with open(jalur, "rb") as berkas:
        data = berkas.read()
    return send_file(BytesIO(data), as_attachment=True, download_name=aman)


@bp.route("/admin/backup/hapus/<nama>", methods=["POST"])
@login_required
def backup_hapus(nama):
    from .arsip import folder_backup, nama_aman

    aman = nama_aman(nama)
    jalur = os.path.join(folder_backup(), aman) if aman else ""
    if aman and os.path.isfile(jalur):
        os.remove(jalur)
        flash("File backup dihapus dari server.", "success")
    else:
        flash("File backup tidak ditemukan.", "error")
    return redirect(url_for("menu.backup"))


@bp.route("/admin/backup/restore", methods=["POST"])
@login_required
def restore():
    from .arsip import MAKS_BYTE, buka

    berkas = request.files.get("berkas")
    if not berkas or not berkas.filename:
        flash("Pilih file backup .eraporbak.", "error")
        return redirect(url_for("menu.backup"))
    blob = berkas.read(MAKS_BYTE + 1)
    if len(blob) > MAKS_BYTE:
        flash("Ukuran file maksimal 100MB.", "error")
        return redirect(url_for("menu.backup"))
    try:
        text = buka(blob, request.form.get("kunci"))
    except ValueError as exc:
        if str(exc) == "ukuran":
            flash("Ukuran file maksimal 100MB.", "error")
        elif str(exc) == "format":
            flash("File bukan backup eRapor yang dikenali.", "error")
        else:
            flash("Kata kunci tidak cocok atau file rusak.", "error")
        return redirect(url_for("menu.backup"))
    if not text.startswith("-- eRapor backup"):
        flash("File bukan backup eRapor yang dikenali.", "error")
        return redirect(url_for("menu.backup"))
    try:
        _jalankan_sql(text)
    except Exception:
        flash("Restore gagal. Data tidak diubah.", "error")
        return redirect(url_for("menu.backup"))
    flash("Backup berhasil dipulihkan.", "success")
    return redirect(url_for("menu.backup"))


def _kembali_dimensi():
    return redirect(
        url_for(
            "menu.dimensi",
            q=request.values.get("q") or "",
            urut=request.values.get("urut") or "urutan",
            arah=request.values.get("arah") or "asc",
            per=request.values.get("per") or 10,
            halaman=request.values.get("halaman") or 1,
            baru=1 if request.form.get("aksi") == "lagi" else None,
        )
    )


def _isi_dimensi():
    kode = (request.form.get("kode") or "").strip().lower()
    nama = (request.form.get("nama") or "").strip()
    indikator = (request.form.get("indikator") or "").strip()
    if not kode or not nama or not indikator:
        flash("Kode, nama dimensi, dan indikator wajib diisi.", "error")
        return None
    if len(kode) > 60:
        flash("Kode dimensi terlalu panjang.", "error")
        return None
    mentah = (request.form.get("urutan") or "").strip()
    if not mentah:
        urutan = 1
    else:
        try:
            urutan = int(mentah)
        except ValueError:
            flash("Urutan tampilan harus angka.", "error")
            return None
    return (
        kode,
        nama,
        indikator,
        1 if request.form.get("kelas_7") else 0,
        1 if request.form.get("kelas_8") else 0,
        1 if request.form.get("kelas_9") else 0,
        urutan,
    )


@bp.route("/admin/dimensi-sikap")
@login_required
def dimensi():
    q = (request.args.get("q") or "").strip()
    per = request.args.get("per", 10, type=int) or 10
    if per not in (10, 25, 50, 100):
        per = 10
    halaman = request.args.get("halaman", 1, type=int) or 1
    if halaman < 1:
        halaman = 1
    urut = "urutan" if request.args.get("urut") != "nama" else "nama"
    arah = "DESC" if request.args.get("arah") == "desc" else "ASC"
    where = ["1=1"]
    params = []
    if q:
        where.append("(nama LIKE %s OR indikator LIKE %s OR kode LIKE %s)")
        like = f"%{q}%"
        params.extend([like, like, like])
    saring = " AND ".join(where)
    order = f"nama {arah}, urutan, id" if urut == "nama" else f"urutan {arah}, id"
    total = query(f"SELECT COUNT(*) AS n FROM dimensi_sikap WHERE {saring}", params, one=True)["n"]
    halaman_max = max((total + per - 1) // per, 1)
    if halaman > halaman_max:
        halaman = halaman_max
    awal = 0 if total == 0 else (halaman - 1) * per + 1
    rows = query(
        f"""SELECT id, kode, nama, indikator, kelas_7, kelas_8, kelas_9, urutan
            FROM dimensi_sikap WHERE {saring}
            ORDER BY {order} LIMIT %s OFFSET %s""",
        (*params, per, (halaman - 1) * per),
    )
    akhir = 0 if total == 0 else awal + len(rows) - 1
    return render_template(
        "menu/dimensi.html",
        rows=rows,
        q=q,
        per=per,
        halaman=halaman,
        total=total,
        awal=awal,
        akhir=akhir,
        urut=urut,
        arah=arah.lower(),
        baru=request.args.get("baru") == "1",
        nav="dimensi",
    )


@bp.route("/admin/dimensi-sikap/simpan", methods=["POST"])
@login_required
def simpan_dimensi():
    data = _isi_dimensi()
    if data is None:
        return _kembali_dimensi()
    item_id = request.form.get("id") or ""
    if item_id and not item_id.isdigit():
        item_id = ""
    bentrok = query(
        "SELECT id FROM dimensi_sikap WHERE kode=%s AND id<>%s",
        (data[0], int(item_id) if item_id else 0),
        one=True,
    )
    if bentrok:
        flash("Kode dimensi sudah digunakan.", "error")
        return _kembali_dimensi()
    if item_id:
        ada = query("SELECT id FROM dimensi_sikap WHERE id=%s", (int(item_id),), one=True)
        if not ada:
            flash("Data tidak ditemukan.", "error")
            return redirect(url_for("menu.dimensi"))
        execute(
            """UPDATE dimensi_sikap
               SET kode=%s, nama=%s, indikator=%s, kelas_7=%s, kelas_8=%s, kelas_9=%s, urutan=%s
               WHERE id=%s""",
            (*data, int(item_id)),
        )
        flash("Dimensi sikap diperbarui.", "success")
        return redirect(url_for("menu.dimensi", q=request.form.get("q") or ""))
    execute(
        """INSERT INTO dimensi_sikap (kode, nama, indikator, kelas_7, kelas_8, kelas_9, urutan)
           VALUES (%s,%s,%s,%s,%s,%s,%s)""",
        data,
    )
    flash("Dimensi sikap disimpan.", "success")
    return _kembali_dimensi()


@bp.route("/admin/dimensi-sikap/tingkat/<int:item_id>", methods=["POST"])
@login_required
def toggle_dimensi(item_id):
    bidang = request.form.get("bidang")
    if bidang in ("kelas_7", "kelas_8", "kelas_9"):
        execute(
            f"UPDATE dimensi_sikap SET {bidang}=%s WHERE id=%s",
            (1 if request.form.get("nyala") else 0, item_id),
        )
    return redirect(
        url_for(
            "menu.dimensi",
            q=request.form.get("q") or "",
            urut=request.form.get("urut") or "urutan",
            arah=request.form.get("arah") or "asc",
            per=request.form.get("per") or 10,
            halaman=request.form.get("halaman") or 1,
        )
    )


@bp.route("/admin/dimensi-sikap/hapus/<int:item_id>", methods=["POST"])
@login_required
def hapus_dimensi(item_id):
    execute("DELETE FROM dimensi_sikap WHERE id=%s", (item_id,))
    flash("Dimensi sikap dihapus.", "success")
    return redirect(url_for("menu.dimensi"))


def _kembali_predikat():
    return redirect(
        url_for(
            "menu.predikat",
            q=request.values.get("q") or "",
            urut=request.values.get("urut") or "urutan",
            arah=request.values.get("arah") or "asc",
            per=request.values.get("per") or 10,
            halaman=request.values.get("halaman") or 1,
            baru=1 if request.form.get("aksi") == "lagi" else None,
        )
    )


def _isi_predikat():
    nama = (request.form.get("nama") or "").strip()
    deskripsi = (request.form.get("deskripsi") or "").strip()
    if not nama:
        flash("Nama predikat wajib diisi.", "error")
        return None
    if len(nama) > 50:
        flash("Nama predikat terlalu panjang.", "error")
        return None
    if len(deskripsi) > 255:
        flash("Awalan deskripsi terlalu panjang.", "error")
        return None
    mentah = (request.form.get("urutan") or "").strip()
    if not mentah:
        urutan = 1
    else:
        try:
            urutan = int(mentah)
        except ValueError:
            flash("Urutan nilai harus angka.", "error")
            return None
    return nama, deskripsi or None, urutan, 1 if request.form.get("aktif") else 0


@bp.route("/admin/predikat-sikap")
@login_required
def predikat():
    q = (request.args.get("q") or "").strip()
    per = request.args.get("per", 10, type=int) or 10
    if per not in (10, 25, 50, 100):
        per = 10
    halaman = request.args.get("halaman", 1, type=int) or 1
    if halaman < 1:
        halaman = 1
    arah = "DESC" if request.args.get("arah") == "desc" else "ASC"
    where = ["1=1"]
    params = []
    if q:
        where.append("(nama LIKE %s OR deskripsi LIKE %s)")
        like = f"%{q}%"
        params.extend([like, like])
    saring = " AND ".join(where)
    total = query(f"SELECT COUNT(*) AS n FROM predikat_sikap WHERE {saring}", params, one=True)["n"]
    halaman_max = max((total + per - 1) // per, 1)
    if halaman > halaman_max:
        halaman = halaman_max
    awal = 0 if total == 0 else (halaman - 1) * per + 1
    rows = query(
        f"""SELECT id, nama, deskripsi, aktif, urutan
            FROM predikat_sikap WHERE {saring}
            ORDER BY urutan {arah}, id LIMIT %s OFFSET %s""",
        (*params, per, (halaman - 1) * per),
    )
    akhir = 0 if total == 0 else awal + len(rows) - 1
    return render_template(
        "menu/predikat.html",
        rows=rows,
        q=q,
        per=per,
        halaman=halaman,
        total=total,
        awal=awal,
        akhir=akhir,
        arah=arah.lower(),
        baru=request.args.get("baru") == "1",
        nav="predikat",
    )


@bp.route("/admin/predikat-sikap/simpan", methods=["POST"])
@login_required
def simpan_predikat():
    data = _isi_predikat()
    if data is None:
        return _kembali_predikat()
    item_id = request.form.get("id") or ""
    if item_id and not item_id.isdigit():
        item_id = ""
    if item_id:
        ada = query("SELECT id FROM predikat_sikap WHERE id=%s", (int(item_id),), one=True)
        if not ada:
            flash("Data tidak ditemukan.", "error")
            return redirect(url_for("menu.predikat"))
        execute(
            "UPDATE predikat_sikap SET nama=%s, deskripsi=%s, urutan=%s, aktif=%s WHERE id=%s",
            (*data, int(item_id)),
        )
        flash("Predikat sikap diperbarui.", "success")
        return redirect(url_for("menu.predikat", q=request.form.get("q") or ""))
    execute(
        "INSERT INTO predikat_sikap (nama, deskripsi, urutan, aktif) VALUES (%s,%s,%s,%s)",
        data,
    )
    flash("Predikat sikap disimpan.", "success")
    return _kembali_predikat()


@bp.route("/admin/predikat-sikap/aktif/<int:item_id>", methods=["POST"])
@login_required
def toggle_predikat(item_id):
    execute(
        "UPDATE predikat_sikap SET aktif=%s WHERE id=%s",
        (1 if request.form.get("nyala") else 0, item_id),
    )
    return redirect(
        url_for(
            "menu.predikat",
            q=request.form.get("q") or "",
            arah=request.form.get("arah") or "asc",
            per=request.form.get("per") or 10,
            halaman=request.form.get("halaman") or 1,
        )
    )


@bp.route("/admin/predikat-sikap/hapus/<int:item_id>", methods=["POST"])
@login_required
def hapus_predikat(item_id):
    execute("DELETE FROM predikat_sikap WHERE id=%s", (item_id,))
    flash("Predikat dihapus.", "success")
    return redirect(url_for("menu.predikat"))


@bp.route("/admin/cetak-tahfiz")
@login_required
def cetak_tahfiz():
    kelas = _kelas_scope() if not _is_admin() else query("SELECT id, nama FROM kelas ORDER BY nama")
    boleh = {str(item["id"]) for item in kelas}
    kelas_id = request.args.get("kelas_id") or ""
    if str(kelas_id) not in boleh:
        kelas_id = ""
    siswa = []
    if kelas_id:
        siswa = query(
            """SELECT s.id, s.nama, s.nis, a.nama AS aspek, n.nilai, n.predikat
               FROM siswa s
               LEFT JOIN nilai_tahfiz n ON n.siswa_id=s.id
               LEFT JOIN aspek_tahfiz a ON a.id=n.aspek_id
               WHERE s.kelas_id=%s AND s.status='aktif'
               ORDER BY s.nama, a.nama""",
            (kelas_id,),
        )
    return render_template("menu/cetak_tahfiz.html", kelas=kelas, kelas_id=str(kelas_id), siswa=siswa, nav="cetaktahfiz")


def _pelajaran_tugas():
    ta_id = _ta_id()
    if not ta_id:
        return []
    sql = """
        SELECT pel.id, k.nama AS kelas_nama, m.nama AS mapel_nama
        FROM pelajaran pel
        JOIN kelas k ON k.id=pel.kelas_id
        JOIN mata_pelajaran m ON m.id=pel.mapel_id
        WHERE pel.tahun_ajaran_id=%s
    """
    params = [ta_id]
    if not _is_admin():
        sql += " AND pel.guru_id=%s"
        params.append(_guru_id() or 0)
    sql += " ORDER BY k.nama, m.nama"
    return query(sql, params)


def _angka_nilai(value):
    teks = "" if value is None else str(value).strip()
    if teks == "":
        return ""
    try:
        angka = float(teks.replace(",", "."))
    except ValueError:
        return None
    if angka < 0 or angka > 100:
        return None
    return round(angka, 2)


@bp.route("/admin/nilai-mapel", methods=["GET"])
@login_required
def nilai_mapel():
    pelajaran = _pelajaran_tugas()
    milik = {str(item["id"]) for item in pelajaran}
    pelajaran_id = request.args.get("pelajaran_id") or (pelajaran[0]["id"] if pelajaran else "")
    if str(pelajaran_id) not in milik:
        pelajaran_id = ""
    siswa = []
    if pelajaran_id:
        siswa = query(
            """SELECT s.id, s.nis, s.nama, nf.nilai AS formatif, ns.nilai AS sumatif
               FROM pelajaran pel
               JOIN siswa s ON s.kelas_id=pel.kelas_id AND s.status='aktif'
               LEFT JOIN nilai_formatif nf ON nf.pelajaran_id=pel.id AND nf.siswa_id=s.id
               LEFT JOIN nilai_sumatif ns ON ns.pelajaran_id=pel.id AND ns.siswa_id=s.id
               WHERE pel.id=%s
               ORDER BY s.nama""",
            (pelajaran_id,),
        )
        for item in siswa:
            for kolom in ("formatif", "sumatif"):
                if item[kolom] is None:
                    item[kolom] = ""
                else:
                    teks = format(item[kolom], "f").rstrip("0").rstrip(".")
                    item[kolom] = teks
    return render_template(
        "menu/nilai_mapel.html",
        pelajaran=pelajaran,
        pelajaran_id=str(pelajaran_id),
        siswa=siswa,
        nav="nilaimapel",
    )


@bp.route("/admin/nilai-mapel/simpan", methods=["POST"])
@login_required
def simpan_nilai_mapel():
    pelajaran_id = request.form.get("pelajaran_id")
    boleh = {str(item["id"]) for item in _pelajaran_tugas()}
    if str(pelajaran_id) not in boleh:
        flash("Pembelajaran ini bukan tugas Anda.", "error")
        return redirect(url_for("menu.nilai_mapel"))
    ids = request.form.getlist("siswa_id")
    siap = []
    for sid in ids:
        for kolom, tabel in (("formatif", "nilai_formatif"), ("sumatif", "nilai_sumatif")):
            angka = _angka_nilai(request.form.get(f"{kolom}_{sid}"))
            if angka is None:
                flash("Nilai harus angka antara 0 dan 100.", "error")
                return redirect(url_for("menu.nilai_mapel", pelajaran_id=pelajaran_id))
            siap.append((tabel, sid, angka))
    with transaction():
        for tabel, sid, angka in siap:
            if angka == "":
                execute(f"DELETE FROM {tabel} WHERE pelajaran_id=%s AND siswa_id=%s", (pelajaran_id, sid))
            else:
                execute(
                    f"""INSERT INTO {tabel} (pelajaran_id, siswa_id, nilai) VALUES (%s,%s,%s)
                        ON DUPLICATE KEY UPDATE nilai=VALUES(nilai)""",
                    (pelajaran_id, sid, angka),
                )
    flash("Nilai mapel disimpan.", "success")
    return redirect(url_for("menu.nilai_mapel", pelajaran_id=pelajaran_id))


@bp.route("/admin/nilai-ekskul", methods=["GET"])
@login_required
def nilai_ekskul():
    if _is_admin():
        ekskul = query("SELECT id, nama FROM ekstrakurikuler ORDER BY nama")
    else:
        ekskul = query("SELECT id, nama FROM ekstrakurikuler WHERE pembina_id=%s ORDER BY nama", (_guru_id() or 0,))
    milik = {str(item["id"]) for item in ekskul}
    ekskul_id = request.args.get("ekskul_id") or (ekskul[0]["id"] if ekskul else "")
    if str(ekskul_id) not in milik:
        ekskul_id = ""
    kelas = query("SELECT id, nama FROM kelas ORDER BY nama")
    kelas_id = request.args.get("kelas_id") or ""
    siswa = []
    if ekskul_id and kelas_id:
        siswa = query(
            """SELECT s.id, s.nama, s.nis, n.nilai
               FROM siswa s
               LEFT JOIN nilai_ekskul n ON n.siswa_id=s.id AND n.ekskul_id=%s
               WHERE s.kelas_id=%s AND s.status='aktif' ORDER BY s.nama""",
            (ekskul_id, kelas_id),
        )
    return render_template(
        "menu/nilai_ekskul.html",
        ekskul=ekskul,
        ekskul_id=str(ekskul_id),
        kelas=kelas,
        kelas_id=str(kelas_id),
        siswa=siswa,
        nav="nilaiekskul",
    )


@bp.route("/admin/nilai-ekskul/simpan", methods=["POST"])
@login_required
def simpan_ekskul():
    ekskul_id = request.form.get("ekskul_id")
    kelas_id = request.form.get("kelas_id")
    if not _is_admin():
        milik = query(
            "SELECT id FROM ekstrakurikuler WHERE id=%s AND pembina_id=%s",
            (ekskul_id, _guru_id() or 0),
            one=True,
        )
        if not milik:
            flash("Ekstrakurikuler ini bukan binaannya.", "error")
            return redirect(url_for("menu.nilai_ekskul"))
    for sid in request.form.getlist("siswa_id"):
        nilai = (request.form.get(f"nilai_{sid}") or "").strip()
        if not nilai:
            continue
        execute(
            """INSERT INTO nilai_ekskul (ekskul_id, siswa_id, nilai) VALUES (%s,%s,%s)
               ON DUPLICATE KEY UPDATE nilai=VALUES(nilai)""",
            (ekskul_id, sid, nilai),
        )
    flash("Nilai ekstrakurikuler disimpan.", "success")
    return redirect(url_for("menu.nilai_ekskul", ekskul_id=ekskul_id, kelas_id=kelas_id))
